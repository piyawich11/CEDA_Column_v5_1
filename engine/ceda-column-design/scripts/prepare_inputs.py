"""Prepare engine inputs from uploads: ETABS exports (+ headerless combo paste) + the CEDA input workbook.

usage:
  python3 -I prepare_inputs.py OUTDIR --seismic S1.xlsx [S2.xlsx ...] [--full-reduced F1.xlsx ...] [--full-unreduced U1.xlsx ...]
                               [--combos LoadCombination.xlsx] [--lap-table DevLength.xlsx [--lap-sheet 40mm]]
                               [--workbook CEDA_Column_Input.xlsx] [--template] [--overload-confirmed]
  v1.3.6: --full-unreduced alone (no reduced-stiffness full model) is used as the primary full model (geometry, βd, basement
  ΔP) and reported as an ASSUMPTION; --combos = Load Combination Definitions file merged into every model; --lap-table = CEDA
  TABLE 11.04 workbook filling the Lap sheet when the input workbook gives none.
  --overload-confirmed: only after the designer confirmed, in this session, every row of the workbook sheet Accept_overload
  (D/C > 1 accepted); without it the rows are kept as accept_overload_pending and NOT applied, and the check exits 2.
  -> OUTDIR/config.json, OUTDIR/<model>.xlsx (normalised ETABS tables), OUTDIR/input_check.json
     (--template also writes OUTDIR/CEDA_Column_Input_prefilled.xlsx for the designer to fill)
then: python3 -I run.py OUTDIR/config.json OUTDIR/OUT ; python3 -I make_outputs.py OUTDIR/OUT --per-mark
JSON in / JSON out; the design rules live in the engine (run.py, design.py, layout.py, colsurf.py), not here.
"""
import os, sys, json, math, statistics
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import etabs_io as E                 # noqa: E402
import input_workbook as INP         # noqa: E402

LAP_KNOWN = INP.LAP_KNOWN

NEEDED = {
    'Element Forces - Columns': 'forces (required)',
    'Load Combination Definitions': 'combo factors – ข้อ 7 check, seismic/gravity split, basement factors',
    'Column Object Connectivity': 'Lu per column-story, story order',
    'Point Object Connectivity': 'XY (label matching between models), elevations',
    'Frame Assignments - Section Properties': 'section of every column-story',
    'Frame Section Property Definitions - Concrete Rectangular': 'section dimensions',
    'Story Forces': 'sway check Q (seismic model)',
    'Story Drifts': 'sway check Q (seismic model)',
}


def _headerless_combos(path):
    """Load Combination Definitions pasted without the TABLE:/header rows (Name, Type, Auto, Load Name, Mode, SF)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows, cur = [], None
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            r = list(r) + [None] * 6
            if r[0] is not None and str(r[0]).upper().startswith('TABLE:'):
                return []
            ld, sf = r[3], r[5]
            if isinstance(ld, str) and isinstance(sf, (int, float)):
                cur = r[0] if r[0] not in (None, '') else cur
                if cur:
                    rows.append({'Name': str(cur), 'Load Name': ld.strip(), 'SF': float(sf)})
    return rows if len(rows) >= 2 else []


def read_upload(path):
    """-> (tables, units, info). Accepts normal ETABS exports and headerless combo definitions."""
    tables, units = E.read_tables(path)
    info = []
    if not tables:
        rows = _headerless_combos(path)
        if rows:
            tables = {'Load Combination Definitions': rows}
            units = {'Load Combination Definitions': {}}
            info.append(f'{os.path.basename(path)}: read as Load Combination Definitions without header ({len(rows)} rows)')
        else:
            info.append(f'{os.path.basename(path)}: no ETABS "TABLE:" block found – ignored')
    return tables, units, info


def merge(parts):
    T, U = {}, {}
    for t, u in parts:
        for k, v in t.items():
            T.setdefault(k, []).extend(v)
            U.setdefault(k, u.get(k, {}))
    return T, U


def write_tables(path, tables, units):
    """write tables back as an ETABS-style workbook (one sheet, TABLE: blocks) for the engine reader."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for name, rows in tables.items():
        if not rows:
            continue
        hdr = []
        for r in rows:
            for k in r:
                if k not in hdr:
                    hdr.append(k)
        ws.append(['TABLE:  ' + name])
        ws.append(hdr)
        u = units.get(name, {})
        ws.append([u.get(h, '') for h in hdr])
        for r in rows:
            ws.append([r.get(h) for h in hdr])
    wb.save(path)


def inventory(tables):
    out = []
    for name, why in NEEDED.items():
        t, real = E.find_table(tables, name)
        out.append(dict(table=name, found=bool(t), rows=len(t) if t else 0, use=why))
    return out


def scan(tables, units, gravity=None):
    """labels, stories, combos and what is already known (sections, Lu, elevations) in one model."""
    gn = gravity or E.GRAV_DEFAULT
    allg = {E._norm(x): g for g, xs in gn.items() for x in xs}
    t, n = E.find_table(tables, 'Element Forces - Columns')
    forces = E.column_forces(tables, units, 'x') if t else []
    labels = sorted({f['label'] for f in forces}, key=lambda s: (len(s), s))
    stories = sorted({f['story'] for f in forces})
    combos = sorted({f['combo'] for f in forces if 'comb' in f['ctype'].lower() or not f['ctype'].strip()})
    cases = sorted({f['combo'] for f in forces} - set(combos))
    fac, kind = E.combo_factors(tables, gn)
    case_names = sorted({c for f in fac.values() for c in f}) or cases
    geo = E.geometry(tables, units)
    secs = E.sections(tables, units)
    order = E.story_order(geo) if geo else []
    from_z = bool(order)
    if not order:
        from run import story_rank
        order = sorted(stories, key=story_rank)
    order = [s for s in order if s in stories] + [s for s in stories if s not in order]
    L, z = {}, {}
    for (lab, st), g in geo.items():
        if 'L' in g:
            L.setdefault(st, []).append(g['L'] / 100)
        if 'zb' in g:
            z[st] = (round(g['zb'] / 100, 3), round(g['zt'] / 100, 3))
    Lu = {st: round(statistics.median(v), 3) for st, v in L.items()}
    sec_of = {}
    for (lab, st), g in geo.items():
        s = g.get('section')
        d = (secs.get(s) or E.section_from_name(s)) if s else None
        if d:
            sec_of.setdefault(lab, set()).add((d['shape'], round(d['t2'] * 10), round(d['t3'] * 10), bool(d.get('assumed'))))
    has_cases = any(allg.get(E._norm(f['combo'])) for f in forces)
    return dict(labels=labels, stories=order, combos=combos, cases=cases, case_names=case_names, defined=sorted(fac), kind=kind,
                Lu=Lu, z=z, sec_of=sec_of, sec_labels=set(sec_of), L_keys={k for k, g in geo.items() if 'L' in g},
                order_from_z=from_z, has_gravity_cases=has_cases, n_forces=len(forces))


def route_models(up, combos=None):
    """v1.3.6: merge a separate Load Combination Definitions upload into every model; without a reduced-stiffness full model
    the unreduced one becomes the primary full model (ASSUMPTION, reported). -> (up, notes)"""
    up, notes = dict(up), []
    if combos:
        new, cur = [], None
        for r in combos[0].get('Load Combination Definitions', []):       # fill the blank names of continuation rows
            nm = E.col(r, 'Name', 'Combo') or cur; cur = nm
            new.append(dict(r, Name=nm))
        mine = {str(r['Name']) for r in new}
        for k in ('seismic', 'full_reduced', 'full_unreduced'):
            if not up.get(k):
                continue
            T, U = dict(up[k][0]), dict(up[k][1])
            old, cur, keep = T.get('Load Combination Definitions', []), None, []
            for r in old:                                 # the separate file wins per combo (no double counting)
                nm = E.col(r, 'Name', 'Combo') or cur; cur = nm
                if str(nm) not in mine:
                    keep.append(dict(r, Name=nm))
            T['Load Combination Definitions'] = keep + new
            U.setdefault('Load Combination Definitions', combos[1].get('Load Combination Definitions', {}))
            up[k] = (T, U)
            if len(keep) != len(old):
                notes.append(f'{k}: {len(old) - len(keep)} ETABS combo rows replaced by the separate combination file')
        notes.append(f'Load combinations: separate file ({len(mine)} combos) used in every model')
    if not up.get('full_reduced') and up.get('full_unreduced'):
        up['full_reduced'], up['full_unreduced'] = up['full_unreduced'], None
        notes.append('ASSUMPTION: no reduced-stiffness full model – the unreduced full model is used for the gravity combos, '
                     'geometry, βd and the basement ΔP (Rev.1 uses the reduced one; upload it when available)')
    elif up.get('full_reduced') and up.get('full_unreduced'):
        notes.append('Full models: reduced stiffness = primary (geometry, βd, basement ΔP); gravity combos enveloped with the unreduced model')
    return up, notes


def prepare(up, tpl):
    """up: {'seismic'|'full_reduced'|'full_unreduced': (tables, units) or None}; tpl: inputs.read_template() dict."""
    single = up.get('full_reduced') is None or up.get('seismic') is None
    base = {'seismic': up.get('seismic') or up.get('full_reduced'), 'full_reduced': up.get('full_reduced') or up.get('seismic'),
            'full_unreduced': up.get('full_unreduced')}
    label_story = set()
    for m in base.values():
        if m:
            label_story |= {(f['label'], f['story']) for f in E.column_forces(m[0], m[1], 'x')}
    eff, rep = {}, []
    for k, m in base.items():
        if m is None:
            eff[k] = None; continue
        T, U, r = INP.apply(m[0], m[1], tpl, label_story)
        eff[k] = (T, U); rep += [f'{k}: {x}' for x in r]
    cfg = INP.to_cfg(tpl)
    gn = cfg['gravity_cases']
    scS, scR = scan(*eff['seismic'], gravity=gn), scan(*eff['full_reduced'], gravity=gn)
    labels = sorted({l for l, s in label_story}, key=lambda s: (len(s), s))
    stories = cfg['story_order'] or (scR['stories'] + [s for s in scS['stories'] if s not in scR['stories']])
    combos = sorted(set(scS['combos']) | set(scR['combos']))
    undefined = [c for c in combos if c not in set(scS['defined']) | set(scR['defined']) and c not in set(cfg['exclude_combos'])]   # v1.3.6: excluded = resolved
    found = dict(labels=labels, stories=stories, combos=combos, undefined=undefined or combos, undefined_only=undefined,
                 cases=scR['case_names'] or scS['case_names'], sec_of={**scS['sec_of'], **scR['sec_of']}, z={**scS['z'], **scR['z']})
    return dict(single=single, eff=eff, merge_rep=rep, cfg_tpl=cfg, scS=scS, scR=scR, labels=labels, stories=stories,
                combos=combos, undefined=undefined, found=found, label_story=label_story)


def build_cfg(prep, workdir):
    paths = {}
    for k, fn in (('seismic', 'seismic.xlsx'), ('full_reduced', 'full_reduced.xlsx'), ('full_unreduced', 'full_unreduced.xlsx')):
        if prep['eff'][k]:
            write_tables(os.path.join(workdir, fn), *prep['eff'][k]); paths[k] = fn
    cfg = json.loads(json.dumps(prep['cfg_tpl']))
    if not cfg['lap_table']:
        cfg['lap_table'] = {str(f): {str(d): v for d, v in t.items()} for f, t in INP.LAP_KNOWN.items()}
    cfg['inputs'] = paths
    if not prep['single']:
        cfg['basement_stories'] = []
    if not cfg['story_order']:
        cfg['story_order'] = prep['stories']
    return cfg


def input_check(prep):
    """the same ✔/✖ list the app shows: every input, its source and whether it is complete."""
    scS, scR, cfg = prep['scS'], prep['scR'], prep['cfg_tpl']
    labels, ls = prep['labels'], prep['label_story']
    sec_ok = {l for l in labels if l in scS['sec_labels'] or l in scR['sec_labels']}
    L_have = scR['L_keys'] | scS['L_keys']
    ov_st = {s for o in cfg['slender']['overrides'] if 'Lu_m' in o and not o['labels'] for s in o['stories']}
    ov_ls = {(l, s) for o in cfg['slender']['overrides'] if 'Lu_m' in o and o['labels'] for l in o['labels'] for s in o['stories']}
    noL = sorted((l, s) for l, s in ls if (l, s) not in L_have and s not in ov_st and (l, s) not in ov_ls)
    bd_ok = scR['has_gravity_cases'] or cfg['slender'].get('beta_d') is not None
    return [
        dict(input='column forces', ok=bool(labels), info=f"{len(labels)} labels, {len(prep['stories'])} stories"),
        dict(input='load combinations', ok=not prep['undefined'], info=', '.join(prep['undefined'][:6])),
        dict(input='column sizes', ok=len(sec_ok) == len(labels), info=', '.join(sorted(set(labels) - sec_ok)[:10])),
        dict(input='Lu / connectivity', ok=not noL, info=f'{len(noL)} column-stories without Lu'),
        dict(input='beta_d', ok=bd_ok, info='from DL/SDL/LL cases' if scR['has_gravity_cases'] else 'Criteria slender.beta_d'),
        dict(input='story order', ok=True, info=' > '.join(prep['stories'][:3]) + ' ... ' + prep['stories'][-1] if prep['stories'] else ''),
    ] + ([dict(input='story ranges (designer)', ok=not INP.check_ranges(cfg['story_ranges'].get('groups'), prep['stories']),
               info='; '.join(INP.check_ranges(cfg['story_ranges'].get('groups'), prep['stories'])) or
               f"{len(cfg['story_ranges']['groups'])} ranges")] if cfg['story_ranges'].get('groups') else []) + ([dict(input='accepted overload (D/C > 1)', ok=bool(prep.get('overload_confirmed')),
               info=f"{len(cfg['accept_overload'])} label-stories in Accept_overload – ask the designer EVERY run, then rerun with "
                    f"--overload-confirmed: " + ', '.join(f'{a} {b} ≤ {c}' for a, b, c in cfg['accept_overload'][:8]))]
         if cfg.get('accept_overload') else [])


def main(a):
    out = a[0]
    os.makedirs(out, exist_ok=True)
    groups, cur = {'seismic': [], 'full_reduced': [], 'full_unreduced': [], 'workbook': [], 'combos': [], 'lap': [], 'lap_sheet': []}, None
    flags = {'--seismic': 'seismic', '--full-reduced': 'full_reduced', '--full-unreduced': 'full_unreduced', '--workbook': 'workbook',
             '--combos': 'combos', '--lap-table': 'lap', '--lap-sheet': 'lap_sheet'}
    confirmed = '--overload-confirmed' in a
    a = [x for x in a if x != '--overload-confirmed']
    for x in a[1:]:
        if x in flags:
            cur = flags[x]
        elif x == '--template':
            cur = None
        elif cur:
            groups[cur].append(x)
    up = {}
    for k in ('seismic', 'full_reduced', 'full_unreduced'):
        parts, info = [], []
        for p in groups[k]:
            t, u, i = read_upload(p); parts.append((t, u)); info += i
        up[k] = merge(parts) if parts else None
        for i in info:
            print('INFO', i)
    cparts = []
    for p in groups['combos']:
        t, u, i = read_upload(p); cparts.append((t, u))
        for x in i:
            print('INFO', x)
    up, route = route_models(up, merge(cparts) if cparts else None)
    for r in route:
        print('NOTE', r)
    empty = dict(criteria={}, stories=[], columns=[], sections=[], combos_rows=[], combo_cases=[], lap_rows=[], exclude=[], approve=[], issues=[])
    tpl = INP.read_template(groups['workbook'][0]) if groups['workbook'] else empty
    if groups['lap'] and not tpl.get('lap_rows'):
        tpl['lap_rows'], note = INP.read_lap_table(groups['lap'][0], groups['lap_sheet'][0] if groups['lap_sheet'] else None)
        print('NOTE lap table:', note)
    prep = prepare(up, tpl)
    prep['merge_rep'] = route + prep['merge_rep']
    prep["overload_confirmed"] = confirmed
    chk = input_check(prep)
    json.dump(dict(check=chk, merge_report=prep['merge_rep'], issues=tpl.get('issues', [])), open(os.path.join(out, 'input_check.json'), 'w'),
              ensure_ascii=False, indent=1)
    if '--template' in a:
        open(os.path.join(out, 'CEDA_Column_Input_prefilled.xlsx'), 'wb').write(INP.make_template(prep['found'], tpl if groups['workbook'] else None))
    cfg = build_cfg(prep, out)
    cfg['inputs'] = {k: v for k, v in cfg['inputs'].items()}
    if cfg.get('accept_overload') and not confirmed:      # CEDA 21 Oct 2026: overload is never applied without asking this run
        cfg['accept_overload_pending'] = cfg.pop('accept_overload')
    json.dump(cfg, open(os.path.join(out, 'config.json'), 'w'), ensure_ascii=False, indent=1)
    for c in chk:
        print(('OK   ' if c['ok'] else 'MISS ') + c['input'] + ('  – ' + c['info'] if c['info'] else ''))
    for r in prep['merge_rep']:
        print('NOTE', r)
    return 0 if all(c['ok'] for c in chk) else 2


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
